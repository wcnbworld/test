import io
import re
from dataclasses import dataclass
from typing import Dict, List, Tuple

import pandas as pd
import streamlit as st
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

REQUIRED_COLUMNS = [
    "公开（公告)号",
    "摘要附图",
    "标题(译)(简体中文)",
    "专利类型",
    "独立权利要求",
    "申请日",
    "[标]当前申请(专利权)人",
    "技术功效",
    "申请号",
    "简单法律状态",
    "摘要(译)(简体中文)",
]

KEEP_CATEGORIES = ["赛车模拟器", "飞行模拟器", "手柄", "渔线轮", "电动扳手"]

APPLICANT_CANONICAL: Dict[str, List[str]] = {
    "成都翼胜": ["成都翼胜科技有限责任公司"],
    "VKB": ["惠州市威科博电子科技有限公司"],
    "VIRPIL": ["UAB VIRPIL"],
    "乌龟海滩": ["沃冶特乌龟海岸公司", "性能设计产品有限责任公司", "乌龟海岸公司", "洛卡特股份有限公司", "性能设计产品有限公司", "英特格姆有限公司"],
    "歌尔": ["歌尔股份有限公司", "歌尔科技有限公司", "歌尔光学科技有限公司", "歌尔微电子股份有限公司", "潍坊歌尔电子有限公司", "青岛歌尔声学科技有限公司", "潍坊歌尔微电子有限公司", "歌尔智能科技有限公司", "青岛歌尔智能传感器有限公司"],
    "飞智": ["上海飞智电子科技有限公司", "上海缕呈科技有限公司"],
    "NACON": ["彼格本因特拉克提夫公司", "BIGBEN CONNECTED", "纳康公司", "讯威科技发展有限公司", "MODELABS GROUP", "BLUETREK TECH"],
    "北通": ["广州市品众电子科技有限公司"],
    "盖世小鸡": ["广州小鸡快跑网络科技有限公司", "深圳闪电鸟网络科技有限公司"],
    "微软": ["微软技术许可有限责任公司", "微软公司", "微软移动设备有限公司", "纽昂斯通讯公司", "金.COM有限公司", "AT&T知识产权二部有限合伙公司"],
    "雷蛇": ["雷蛇(亚太)私人有限公司", "瑞瑟美国公司"],
    "莱仕达&东莞星辰": ["东莞市星辰互动电子科技有限公司", "深圳市莱仕达电子科技有限公司"],
    "罗技": ["罗技欧洲股份有限公司", "罗技电子股份有限公司", "3D连接股份有限公司", "LOGITECH SA", "LOGITECH INT"],
    "速魔": ["深圳市速魔科技有限公司", "李金全", "速模实业(深圳)有限公司"],
    "图马思特": ["基利摩股份有限公司"],
    "SIMUCUBE": ["格兰尼特设备有限公司"],
    "FANATEC": ["可赛尔内存股份有限公司", "铁堡发明有限公司", "恩德游戏设备有限公司", "埃尔加托艾迪斯普雷有限公司", "CORSAIR COMPONENTS LTD", "THE BURGESS BEDDING CO LTD"],
    "索尼": ["索尼集团公司", "索尼互动娱乐有限责任公司", "索尼半导体解决方案公司", "索尼爱立信移动通讯股份有限公司(瑞典)", "索尼电子有限公司", "新力欧洲股份有限公司"],
    "ASETEK": ["阿塞泰克丹麥公司", "阿塞泰克公司"],
    "深圳景创": ["深圳市景创科技电子股份有限公司", "深圳景创智航技术有限公司", "深圳景创智造有限公司"],
    "恩速": ["恩速(上海)电子科技有限公司"],
    "卡妙思": ["深圳市卡妙思电子科技有限公司", "深圳市卡斯妙汽车科技有限公司", "深圳市卡妙思赛车有限公司", "深圳市风翔汽车科技有限公司"],
    "达实智控": ["深圳市达实智控科技股份有限公司"],
    "易速马": ["深圳易速马网络科技有限公司"],
    "八位堂": ["深圳市百思度科技有限公司", "深圳市壹位堂科技有限公司", "深圳市八位堂科技有限公司", "深圳市吉窝窝科技有限公司"],
}

APPLICANT_LOOKUP = {
    re.sub(r"\s+", "", alias): canon
    for canon, aliases in APPLICANT_CANONICAL.items()
    for alias in aliases
}


@dataclass
class PatentRow:
    publication_no: str
    image: str
    title: str
    patent_type: str
    claim: str
    apply_date: str
    applicant: str
    effect: str
    app_no: str
    legal_status: str
    abstract: str
    category: str


def normalize_text(value: str) -> str:
    return re.sub(r"\s+", "", str(value or "")).strip()


def canonicalize_applicant(raw: str) -> str:
    compact = normalize_text(raw)
    return APPLICANT_LOOKUP.get(compact, str(raw or "").strip())


def classify_patent(title: str, claim: str, effect: str, abstract: str) -> str:
    text = " ".join([str(title or ""), str(claim or ""), str(effect or ""), str(abstract or "")]).lower()

    if any(k in text for k in ["赛车", "方向盘", "踏板", "力反馈", "sim racing", "racing simulator"]):
        return "赛车模拟器"
    if any(k in text for k in ["飞行", "flight", "摇杆", "节流阀", "油门杆", "舵"]):
        return "飞行模拟器"
    if any(k in text for k in ["游戏手柄", "控制器", "controller", "gamepad", "joystick"]):
        return "手柄"
    if any(k in text for k in ["渔线轮", "卷线器", "fishing reel"]):
        return "渔线轮"
    if any(k in text for k in ["电动扳手", "电动螺丝刀", "electric wrench", "electric screwdriver"]):
        return "电动扳手"
    return "其他"


def apply_table_border(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        elem = OxmlElement(f"w:{side}")
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "8")
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), "000000")
        borders.append(elem)
    tbl_pr.append(borders)


def patent_type_bucket(pt: str) -> str:
    t = str(pt or "")
    if "外观" in t:
        return "外观"
    if "发明" in t or "实用新型" in t:
        return "发明/实用新型"
    return "其他"


def build_rows(df: pd.DataFrame) -> List[PatentRow]:
    rows: List[PatentRow] = []
    for _, r in df.iterrows():
        title = r.get("标题(译)(简体中文)", "")
        claim = r.get("独立权利要求", "")
        effect = r.get("技术功效", "")
        abstract = r.get("摘要(译)(简体中文)", "")
        rows.append(
            PatentRow(
                publication_no=str(r.get("公开（公告)号", "") or ""),
                image="",
                title=str(title or ""),
                patent_type=str(r.get("专利类型", "") or ""),
                claim=str(claim or ""),
                apply_date=str(r.get("申请日", "") or ""),
                applicant=canonicalize_applicant(r.get("[标]当前申请(专利权)人", "")),
                effect=str(effect or ""),
                app_no=str(r.get("申请号", "") or ""),
                legal_status=str(r.get("简单法律状态", "") or ""),
                abstract=str(abstract or ""),
                category=classify_patent(title, claim, effect, abstract),
            )
        )
    return rows


def add_table_b(doc: Document, rows: List[PatentRow], applicant: str):
    if not rows:
        return
    doc.add_heading(f"{applicant} - 表格B（外观专利）", level=2)
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "申请信息"
    table.rows[0].cells[1].text = "专利名称"
    table.rows[0].cells[2].text = "专利图片"
    for r in rows:
        row = table.add_row().cells
        row[0].text = f"公开（公告)号：{r.publication_no}\n简单法律状态：{r.legal_status}\n申请日：{r.apply_date}\n[标]当前申请(专利权)人：{r.applicant}"
        row[1].text = r.title
        row[2].text = ""
    apply_table_border(table)


def add_table_c(doc: Document, rows: List[PatentRow], applicant: str):
    if not rows:
        return
    doc.add_heading(f"{applicant} - 表格C（发明/授权发明/实用新型）", level=2)
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "申请信息"
    table.rows[0].cells[1].text = "专利图片"
    table.rows[0].cells[2].text = "专利方案"
    for r in rows:
        row = table.add_row().cells
        row[0].text = f"公开（公告)号：{r.publication_no}\n简单法律状态：{r.legal_status}\n标题(译)(简体中文)：{r.title}\n申请日：{r.apply_date}\n[标]当前申请(专利权)人：{r.applicant}"
        row[1].text = ""
        row[2].text = r.effect
    apply_table_border(table)


def add_table_d(doc: Document, rows: List[PatentRow]):
    if not rows:
        return
    doc.add_heading("表格D（技术分类=其他）", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "申请信息"
    table.rows[0].cells[1].text = "专利名称"
    for r in rows:
        row = table.add_row().cells
        row[0].text = f"公开（公告)号：{r.publication_no}\n简单法律状态：{r.legal_status}\n申请日：{r.apply_date}\n[标]当前申请(专利权)人：{r.applicant}"
        row[1].text = r.title
    apply_table_border(table)


def generate_docs(rows: List[PatentRow]) -> Tuple[bytes, bytes, pd.DataFrame]:
    kept = [r for r in rows if r.category in KEEP_CATEGORIES]
    other = [r for r in rows if r.category == "其他"]

    summary_doc = Document()
    summary_doc.add_heading("专利分类结果（按申请人 -> 专利类型）", level=1)

    grouped: Dict[str, List[PatentRow]] = {}
    for r in kept:
        grouped.setdefault(r.applicant, []).append(r)

    for applicant in sorted(grouped.keys()):
        applicant_rows = grouped[applicant]
        design_rows = [r for r in applicant_rows if patent_type_bucket(r.patent_type) == "外观"]
        invent_rows = [r for r in applicant_rows if patent_type_bucket(r.patent_type) == "发明/实用新型"]

        if design_rows:
            add_table_b(summary_doc, design_rows, applicant)
        if invent_rows:
            add_table_c(summary_doc, invent_rows, applicant)

    other_doc = Document()
    add_table_d(other_doc, other)

    out1 = io.BytesIO()
    summary_doc.save(out1)
    out2 = io.BytesIO()
    other_doc.save(out2)

    data = pd.DataFrame(
        [
            {
                "公开（公告)号": r.publication_no,
                "申请人(归一化)": r.applicant,
                "专利类型": r.patent_type,
                "技术分类": r.category,
                "标题": r.title,
            }
            for r in rows
        ]
    )
    return out1.getvalue(), out2.getvalue(), data


def load_dataframe(uploaded_file) -> pd.DataFrame:
    name = uploaded_file.name.lower()
    if name.endswith(".csv"):
        return pd.read_csv(uploaded_file)
    return pd.read_excel(uploaded_file)


def main():
    st.set_page_config(page_title="专利表格分类工具", layout="wide")
    left, right = st.columns(2)

    with left:
        st.header("输入窗口")
        uploaded = st.file_uploader("上传表格A（xlsx/xls/csv）", type=["xlsx", "xls", "csv"])
        st.caption("系统将自动按申请人与专利类型分组，并生成两个 Word 文档。")

    with right:
        st.header("输出窗口")
        if uploaded is None:
            st.info("请先上传表格A。")
            return
        df = load_dataframe(uploaded)
        missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
        if missing:
            st.error(f"缺少必要列：{', '.join(missing)}")
            return

        rows = build_rows(df)
        doc_main, doc_other, result_df = generate_docs(rows)

        st.success(f"处理完成：共 {len(rows)} 条专利。")
        st.dataframe(result_df, use_container_width=True)

        st.download_button(
            "下载主文档（表格B/C）",
            data=doc_main,
            file_name="专利分类结果_表格BC.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        st.download_button(
            "下载其他分类文档（表格D）",
            data=doc_other,
            file_name="专利分类结果_表格D.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )


if __name__ == "__main__":
    main()
